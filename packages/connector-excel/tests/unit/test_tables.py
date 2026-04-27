"""Unit tests for table detection."""

from pathlib import Path

import openpyxl

from excel_connector.models import TableConfidence
from excel_connector.tables import detect_tables


def _ws(path: Path, name: str | None = None):
    wb = openpyxl.load_workbook(path, data_only=False)
    return wb[name] if name else wb.active


def test_listobject_detected_with_certain_confidence(simple_xlsx: Path) -> None:
    ws = _ws(simple_xlsx)
    tables = detect_tables(ws)
    assert len(tables) == 1
    table = tables[0]
    assert table.confidence == TableConfidence.CERTAIN
    assert table.name == "SalesTable"
    assert table.row_count == 3
    assert [c.name for c in table.columns] == ["id", "name", "amount"]


def test_heuristic_detection_when_no_listobject(heuristic_xlsx: Path) -> None:
    ws = _ws(heuristic_xlsx)
    tables = detect_tables(ws)
    assert len(tables) == 1
    table = tables[0]
    assert table.confidence == TableConfidence.INFERRED
    assert table.row_count == 5
    assert {c.name for c in table.columns} == {"sku", "qty", "price"}


def test_formulas_collected_from_table_cells(multi_sheet_xlsx: Path) -> None:
    wb = openpyxl.load_workbook(multi_sheet_xlsx, data_only=False)
    orders_ws = wb["Orders"]
    tables = detect_tables(orders_ws)
    assert len(tables) == 1
    formulas = tables[0].formulas
    assert len(formulas) == 2
    assert all(f.is_lookup for f in formulas)
    assert all("Customers" in f.precedent_sheets for f in formulas)
