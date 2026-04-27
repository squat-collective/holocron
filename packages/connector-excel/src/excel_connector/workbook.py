"""Top-level scan entry point — orchestrates extractors and produces a ScanResult."""

from pathlib import Path

import openpyxl

from excel_connector.actors import extract_actors
from excel_connector.external_links import detect_external_links
from excel_connector.metadata import extract_metadata
from excel_connector.models import (
    DetectedSheet,
    ScanResult,
)
from excel_connector.tables import detect_tables


def scan_workbook(path: str | Path) -> ScanResult:
    """Scan an Excel workbook and return everything we can detect.

    Args:
        path: Path to a .xlsx / .xlsm / .xltx file.

    Returns:
        A populated ScanResult. The caller is responsible for mapping it to
        Holocron API payloads (see excel_connector.mapping).
    """
    abs_path = str(Path(path).resolve())
    file_name = Path(path).name

    # data_only=False so formulas come back as strings (=VLOOKUP(...)), not their cached value.
    wb = openpyxl.load_workbook(abs_path, data_only=False, read_only=False)

    sheets: list[DetectedSheet] = []
    for ws in wb.worksheets:
        tables = detect_tables(ws)
        sheets.append(
            DetectedSheet(
                name=ws.title,
                visible=ws.sheet_state == "visible",
                tables=tables,
            )
        )

    external_links = detect_external_links(wb)

    # Augment external_links with which sheets reference them (from formula precedents)
    if external_links:
        ext_targets = {link.target_path: link for link in external_links}
        for sheet in sheets:
            for table in sheet.tables:
                for formula in table.formulas:
                    for ext_file in formula.precedent_external_files:
                        # Match by basename (formula refs are typically just the filename)
                        for target_path, link in ext_targets.items():
                            if Path(target_path).name == ext_file or target_path == ext_file:
                                if sheet.name not in link.referenced_from_sheets:
                                    link.referenced_from_sheets.append(sheet.name)

    actors = extract_actors(wb)
    metadata = extract_metadata(wb)

    wb.close()

    return ScanResult(
        file_path=abs_path,
        file_name=file_name,
        workbook_metadata=metadata,
        sheets=sheets,
        external_links=external_links,
        actors=actors,
    )
